"""Audit salaries against OEDE/INDEC and import aggregate census/SNIC context.

Inputs are official files downloaded and retained outside the public repository.
python scripts_import_municipal_context.py --input-dir PATH --research-dir PATH
"""
import argparse,csv,hashlib,json,math,statistics
from datetime import datetime
from pathlib import Path
import openpyxl
import xlrd

ROOT=Path(__file__).resolve().parent

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def numeric(v):return isinstance(v,(int,float)) and math.isfinite(v)
def as_number(v):return None if v in ('NA','',None) else float(str(v).replace(',','.'))
def read_book(path):return openpyxl.load_workbook(path,data_only=True,read_only=True)

def build(folder,research):
    data=json.loads((ROOT/'municipios/data/dashboard.json').read_text(encoding='utf-8'))
    municipalities={m['id']:m for m in data['municipalities']}
    sources={r['file']:r for r in json.loads((folder/'downloads.json').read_text(encoding='utf-8')) if 'error' not in r}
    for name,r in sources.items():
        if sha(folder/name)!=r['sha256']:raise ValueError('Source changed: '+name)
    # INDEC national general index, not monthly percentage changes or a regional index.
    ipc_path=research.parent/'municipios-auditoria-20260907/raw/ipc-oficial.xls'
    ipc_book=xlrd.open_workbook(ipc_path);sheet=ipc_book.sheet_by_name('Índices IPC Cobertura Nacional')
    assert sheet.cell_value(5,0)=='Total nacional' and sheet.cell_value(9,0)=='Nivel general'
    ipc={xlrd.xldate_as_datetime(sheet.cell_value(5,c),ipc_book.datemode).strftime('%Y-%m'):sheet.cell_value(9,c) for c in range(1,sheet.ncols) if sheet.cell_type(5,c)==xlrd.XL_CELL_DATE}
    verified_ipc={r['period']:float(r['ipc_index']) for r in csv.DictReader((research/'ipc_indec_verificado.csv').open(encoding='utf-8'))}
    assert all(abs(ipc[k]-v)<1e-8 for k,v in verified_ipc.items()),'IPC differs from INDEC'
    sources['ipc-oficial.xls']={'file':'ipc-oficial.xls','url':'https://www.indec.gob.ar/ftp/cuadros/economia/sh_ipc_08_26.xls','sha256':sha(ipc_path),'bytes':ipc_path.stat().st_size}
    oede=read_book(folder/'oede.xlsx');raw_rows={};series={k:{} for k in municipalities}
    seed=list(csv.DictReader((research/'datos_de_entrada/empleo_salarios_oede_muestra.csv').open(encoding='utf-8-sig')))
    for r in seed:
        if r['metric']!='wage_ars':continue
        key=(r['source_sheet'],int(r['source_row']))
        if key not in raw_rows:
            ws=oede[key[0]];row=next(ws.iter_rows(min_row=key[1],max_row=key[1],values_only=True));header=next(ws.iter_rows(min_row=2,max_row=2,values_only=True))
            assert row[2]=='BUENOS AIRES'
            raw_rows[key]={p.strftime('%Y-%m'):v for p,v in zip(header,row) if isinstance(p,datetime)}
        value=raw_rows[key][r['period']]
        expected=float(r['value']) if r['value'] else None
        assert (not numeric(value) and expected is None) or abs(value-expected)<.00001,('OEDE source mismatch',r['municipality_id'],r['period'])
        series[r['municipality_id']][r['period']]=value if numeric(value) else None
    audit=[];max_error=0;checked=0;records={}
    for ident,m in municipalities.items():
        wages=series[ident]
        assert len(wages)==len(m['employment'])
        for period,jobs,real in m['employment']:
            original=wages[period]
            if original is None:assert real is None
            else:
                rebuilt=original*ipc['2026-07']/ipc[period];error=abs(rebuilt-real);max_error=max(max_error,error);checked+=1
                assert error<=.00501,(ident,period,error)
        annual={}
        for year in (2023,2024,2025):
            months=[p for p in wages if p.startswith(str(year))];assert len(months)==12
            assert all(numeric(wages[p]) for p in months)
            nominal=statistics.mean(wages[p] for p in months);real=statistics.mean(wages[p]*ipc['2026-07']/ipc[p] for p in months)
            assert abs(real-m[f'salario_real_promedio_{year}_ars_jul26'])<.00001
            annual[str(year)]={'nominal':nominal,'real':real}
        references=sorted({(r['source_sheet'],int(r['source_row'])) for r in seed if r['municipality_id']==ident and r['metric']=='wage_ars'})
        wage={'annual':annual,'months':{p:{'nominal':wages[p],'real':wages[p]*ipc['2026-07']/ipc[p]} for p in ['2025-11','2025-12']},'references':[{'sheet':s,'row':r} for s,r in references]}
        records[ident]={'id':ident,'municipality':m['municipio'],'wage':wage}
    # Denominators are kept with each statistic (persons in private dwellings / households).
    for name in ['health.xlsx','crowding.xlsx']:
        seen=set();ws=read_book(folder/name).active
        for line,row in enumerate(ws.iter_rows(values_only=True),1):
            try:ident='06'+str(int(row[0])).zfill(3)
            except (ValueError,TypeError):continue
            assert ident in municipalities and ident not in seen;seen.add(ident)
            if name=='health.xlsx':
                total,insured,plan,none=row[2:6];assert total==insured+plan+none
                records[ident]['health']={'year':2022,'populationPrivateDwellings':total,'socialInsuranceOrPrivate':insured,'statePlan':plan,'withoutCoverage':none,'withoutCoveragePct':100*none/total,'sourceRow':line,'sourceSheet':ws.title}
            else:
                total,low,middle,high=row[2:6];assert abs(low+middle+high-100)<.001
                assert total==municipalities[ident]['hogares_2022']
                records[ident]['crowding']={'year':2022,'households':total,'over3PersonsPerRoomPct':high,'sourceRow':line,'sourceSheet':ws.title}
        assert seen==set(municipalities),(name,len(seen))
    rows=[]
    # The published CSV uses these codes, but explicitly names the districts.
    # Retain original identifiers and rows as evidence; 06999 is unassigned.
    snic_codes={'06058':('06658','Quilmes'),'06217':('06218','Chascomús')}
    with (folder/'crime.csv').open(encoding='utf-8-sig',newline='') as stream:
        for line,r in enumerate(csv.DictReader(stream,delimiter=';'),2):
            if r['provincia_id']!='06' or r['anio'] not in ['2024','2025'] or r['codigo_delito_snic_id'] not in ['1','15','17','19']:continue
            ident=r['departamento_id']
            if ident in snic_codes:
                ident,expected_name=snic_codes[ident];assert r['departamento_nombre']==expected_name
            if ident in municipalities:rows.append({**r,'municipalityId':ident,'sourceLine':line})
    assert len(rows)==135*2*4
    for ident,record in records.items():
        record['crime']={}
        for year in ['2024','2025']:
            chosen={r['codigo_delito_snic_id']:r for r in rows if r['municipalityId']==ident and r['anio']==year};assert len(chosen)==4
            def count(code,field):return as_number(chosen[code][field])
            for code in chosen:
                assert count(code,'cantidad_hechos') is not None
            record['crime'][year]={'homicideVictims':count('1','cantidad_victimas'),'homicideRate':count('1','tasa_victimas'),
                'robberies':count('15','cantidad_hechos')+count('17','cantidad_hechos'),'robberyRate':count('15','tasa_hechos')+count('17','tasa_hechos'),
                'thefts':count('19','cantidad_hechos'),'theftRate':count('19','tasa_hechos'),'sourceRows':chosen}
    credit=read_book(folder/'household-credit.xlsx')['4.1.3'];headers=list(credit.iter_rows(min_row=5,max_row=5,values_only=True))[0];ba=list(credit.iter_rows(min_row=6,max_row=6,values_only=True))[0]
    last=max((p,v) for p,v in zip(headers,ba) if isinstance(p,datetime) and numeric(v))
    result={'version':1,'verified':'2026-09-09','priceBase':'2026-07','sources':sources,
        'salaryAudit':{'municipalities':135,'monthlyObservations':checked,'maxRoundingDifferenceARS':max_error,'formula':'salary_month * IPC_2026_07 / IPC_month; annual = mean(12 months)','result':'All monthly and annual values reconcile to OEDE and INDEC.'},
        'householdDebt':{'geographicLevel':'province','municipalAvailable':False,'provincialPeriod':last[0].strftime('%Y-%m'),'buenosAiresAdultsWithCreditPct':last[1],'note':'Public BCRA credit series is provincial; do not assign it to municipalities or equate local bank loans with household debt.'},
        'municipalities':list(records.values())}
    (ROOT/'municipios/data/community_verified.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'salaryAudit':result['salaryAudit'],'lasHeras':records['06329'],'householdDebt':result['householdDebt']},ensure_ascii=False))
if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--input-dir',type=Path,required=True);parser.add_argument('--research-dir',type=Path,required=True);args=parser.parse_args();build(args.input_dir,args.research_dir)
