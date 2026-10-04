import csv,hashlib,json,unittest
from pathlib import Path
import openpyxl
from scripts_export_management_reports import ReportData
ROOT=Path(__file__).resolve().parents[1]
class SeptemberRON(unittest.TestCase):
 def test_official_totals_and_shared_national_supplement(self):
  d=json.loads((ROOT/'data/ron_latest.json').read_text(encoding='utf8'))
  self.assertEqual(d['period'],'2026-09');self.assertEqual(d['real_available_through'],'2026-08')
  self.assertEqual(len(d['rows']),24)
  p=ROOT/d['source']['path'];self.assertEqual(hashlib.sha256(p.read_bytes()).hexdigest(),d['source']['sha256'])
  w=openpyxl.load_workbook(p,data_only=True)
  names={'C.A.B.A':'CABA','C.A.B.A.':'CABA','Sgo. Del Estero':'Santiago del Estero','Sgo. del Estero':'Santiago del Estero','Tierra Del Fuego':'Tierra del Fuego'}
  for s,col in [('septiembre','month_nominal'),('CONS','ytd_nominal')]:
   sheet=w[s];values={names.get(sheet.cell(i,1).value,sheet.cell(i,1).value):sheet.cell(i,29).value for i in range(1,79)}
   for r in d['rows']:self.assertAlmostEqual(r[col],values[r['province']],delta=.1)
  n=json.loads((ROOT/'nacion/data/gestion.json').read_text(encoding='utf8'))
  self.assertEqual(n['provinces']['latest'],d)
 def test_no_invented_september_real_change(self):
  data=ReportData();p=data.transfers('Buenos Aires')
  self.assertEqual(p['through'],'2026-09');self.assertIsNone(p['real_pct'])
  self.assertAlmostEqual(p['current'],13266974.425482867,places=5)
  self.assertIsNotNone(data.transfers('Buenos Aires','2026-08')['real_pct'])
if __name__=='__main__':unittest.main()
